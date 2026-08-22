from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EMPTY = (None, "")
KEY_WORDS = re.compile(r"(^id$|\bid\b|kode|code|nomor|number|no\.?$|name|nama)", re.I)
SENSITIVE_HINT = re.compile(r"password|secret|token|credential", re.I)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def kind(value: Any) -> str:
    if value in EMPTY:
        return "empty"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    return "text"


def normalized(value: Any) -> str:
    if value in EMPTY:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def header_score(values: list[Any]) -> float:
    nonempty = [v for v in values if v not in EMPTY]
    if not nonempty:
        return -1.0
    text = sum(isinstance(v, str) and not str(v).startswith("=") for v in nonempty)
    semantic = sum(bool(re.search(r"[A-Za-zÀ-ÿ]", str(v))) for v in nonempty)
    unique = len({normalized(v) for v in nonempty})
    return len(nonempty) + 2 * text / len(nonempty) + semantic / len(nonempty) + unique / len(nonempty)


def header_row(ws, scan_rows: int = 20) -> int:
    candidates = []
    for row_no in range(1, min(ws.max_row, scan_rows) + 1):
        values = [ws.cell(row_no, c).value for c in range(1, ws.max_column + 1)]
        candidates.append((header_score(values), row_no))
    return max(candidates or [(0.0, 1)])[1]


def headers(ws, row_no: int) -> list[str]:
    seen: Counter[str] = Counter()
    result = []
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(row_no, col).value
        base = re.sub(r"\s+", " ", str(raw).strip()) if raw not in EMPTY else f"UNNAMED_{col}"
        seen[base] += 1
        result.append(base if seen[base] == 1 else f"{base}__{seen[base]}")
    return result


def profile_sheet(ws) -> dict[str, Any]:
    hrow = header_row(ws)
    names = headers(ws, hrow)
    data_rows = range(hrow + 1, ws.max_row + 1)
    columns = []
    key_candidates = []
    row_fingerprints = Counter()

    for r in data_rows:
        vals = [normalized(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if any(vals):
            row_fingerprints[tuple(vals)] += 1

    for col, name in enumerate(names, 1):
        vals = [ws.cell(r, col).value for r in data_rows]
        nonempty = [v for v in vals if v not in EMPTY]
        counts = Counter(kind(v) for v in vals)
        norm = [normalized(v) for v in nonempty]
        distinct = len(set(norm))
        duplicates = len(norm) - distinct
        candidate = bool(nonempty and not duplicates and len(nonempty) == len(vals) and KEY_WORDS.search(name))
        if candidate:
            key_candidates.append(name)
        columns.append({
            "name": "REDACTED" if SENSITIVE_HINT.search(name) else name,
            "types": dict(counts),
            "nonempty": len(nonempty),
            "empty": len(vals) - len(nonempty),
            "distinct": distinct,
            "duplicate_values": duplicates,
            "candidate_key": candidate,
        })

    formula_count = sum(
        1 for row in ws.iter_rows() for cell in row
        if cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))
    )
    duplicate_groups = sum(1 for count in row_fingerprints.values() if count > 1)
    duplicate_rows = sum(count - 1 for count in row_fingerprints.values() if count > 1)
    return {
        "name": ws.title,
        "state": ws.sheet_state,
        "rows": ws.max_row,
        "columns_count": ws.max_column,
        "header_row": hrow,
        "headers": names,
        "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
        "formula_count": formula_count,
        "duplicate_row_groups": duplicate_groups,
        "duplicate_rows_extra": duplicate_rows,
        "candidate_keys": key_candidates,
        "columns": columns,
    }


def audit(path: Path) -> dict[str, Any]:
    wb = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        return {
            "file": path.name,
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
            "sheet_count": len(wb.sheetnames),
            "defined_names": sorted(str(k) for k in wb.defined_names),
            "sheets": [profile_sheet(wb[name]) for name in wb.sheetnames],
        }
    finally:
        wb.close()


def self_check() -> None:
    assert normalized(" A  B ") == "a b"
    assert kind("=SUM(A1:A2)") == "formula"
    assert kind(None) == "empty"


def main() -> None:
    parser = argparse.ArgumentParser(description="Aggregate-only XLSX audit; never writes source files")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    self_check()
    files = sorted(args.source.glob("*.xlsx"))
    result = {"source": str(args.source), "workbook_count": len(files), "workbooks": [audit(p) for p in files]}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps({"workbooks": len(files), "sheets": sum(w["sheet_count"] for w in result["workbooks"]), "output": str(args.output)}))


if __name__ == "__main__":
    main()
