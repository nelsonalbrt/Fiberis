from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Sumber data excel"
PROFILE = json.loads((ROOT / "audit/workbook_profile.json").read_text(encoding="utf-8"))
GENERIC = {"no", "nomor", "keterangan", "ket", "status", "nama", "unnamed", "port", "kapasitas", "action", "remark", "remarks"}


def canon(name: str) -> str:
    value = re.sub(r"__\d+$", "", name.casefold())
    value = re.sub(r"[^a-z0-9]+", " ", value).strip()
    return re.sub(r"\s+", " ", value)


def useful(name: str) -> bool:
    value = canon(name)
    return bool(value and not value.startswith("unnamed") and value not in GENERIC and len(value) >= 3)


def norm(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


columns = defaultdict(list)
for wb_meta in PROFILE["workbooks"]:
    path = SOURCE / wb_meta["file"]
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for sheet_meta in wb_meta["sheets"]:
            ws = wb[sheet_meta["name"]]
            selected = {
                index: (header, set())
                for index, header in enumerate(sheet_meta["headers"])
                if useful(header)
            }
            if not selected:
                continue
            for row in ws.iter_rows(
                min_row=sheet_meta["header_row"] + 1,
                max_row=sheet_meta["rows"],
                values_only=True,
            ):
                for index, (_, values) in selected.items():
                    if index < len(row):
                        value = norm(row[index])
                        if value:
                            values.add(value)
            for header, values in selected.values():
                if len(values) < 5:
                    continue
                columns[canon(header)].append({
                    "workbook": wb_meta["file"], "sheet": ws.title, "column": header, "values": values,
                })
    finally:
        wb.close()

candidates = []
for name, refs in columns.items():
    for i, left in enumerate(refs):
        for right in refs[i + 1:]:
            if left["workbook"] == right["workbook"] and left["sheet"] == right["sheet"]:
                continue
            intersection = len(left["values"] & right["values"])
            if intersection < 3:
                continue
            smaller = min(len(left["values"]), len(right["values"]))
            union = len(left["values"] | right["values"])
            containment = intersection / smaller
            jaccard = intersection / union
            if containment < 0.5:
                continue
            candidates.append({
                "canonical_header": name,
                "left": {k: v for k, v in left.items() if k != "values"},
                "right": {k: v for k, v in right.items() if k != "values"},
                "left_distinct": len(left["values"]),
                "right_distinct": len(right["values"]),
                "intersection": intersection,
                "containment": round(containment, 4),
                "jaccard": round(jaccard, 4),
                "confidence": "INFERRED",
            })

candidates.sort(key=lambda x: (x["containment"], x["intersection"], x["jaccard"]), reverse=True)
output = {"candidate_count": len(candidates), "candidates": candidates[:200]}
(ROOT / "audit/relation_candidates.json").write_text(json.dumps(output, indent=2, ensure_ascii=False), encoding="utf-8")
print(json.dumps({"candidate_count": len(candidates), "written": min(200, len(candidates))}))
