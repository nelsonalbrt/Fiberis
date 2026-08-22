from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EMPTY = (None, "")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def value_text(value: Any) -> str:
    if value in EMPTY:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def choose_header(ws, limit: int = 20) -> int:
    scored = []
    for row in range(1, min(ws.max_row, limit) + 1):
        values = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
        nonempty = [value for value in values if value not in EMPTY]
        text = sum(isinstance(value, str) for value in nonempty)
        scored.append((len(nonempty) + text, row))
    return max(scored or [(0, 1)])[1]


def sheet_preview(ws) -> dict[str, Any]:
    header = choose_header(ws)
    seen: Counter[tuple[str, ...]] = Counter()
    blank = duplicate = review = 0
    issues: Counter[str] = Counter()
    examples = []
    for row in range(header + 1, ws.max_row + 1):
        values = tuple(value_text(ws.cell(row, col).value) for col in range(1, ws.max_column + 1))
        if not any(values):
            blank += 1
            issues["BLANK_ROW"] += 1
            continue
        seen[values] += 1
        if seen[values] > 1:
            duplicate += 1
            issues["EXACT_DUPLICATE"] += 1
            if len(examples) < 20:
                examples.append({"row": row, "status": "DUPLICATE", "issue": "EXACT_DUPLICATE"})
            continue
        review += 1
    return {
        "name": ws.title,
        "state": ws.sheet_state,
        "header_row": header,
        "rows": max(0, ws.max_row - header),
        "blank": blank,
        "duplicate": duplicate,
        "review": review,
        "issues": dict(issues),
        "examples": examples,
        "publishable": False,
    }


def preview(source: Path) -> dict[str, Any]:
    workbooks = []
    total = Counter(workbooks=0, sheets=0, rows=0, blank=0, duplicate=0, review=0, blocked=0)
    for path in sorted(source.glob("*.xlsx")):
        try:
            before = sha256(path)
        except PermissionError:
            total["blocked"] += 1
            workbooks.append({"file": path.name, "status": "BLOCKED", "issue": "FILE_LOCKED", "sheets": []})
            continue
        wb = load_workbook(path, read_only=False, data_only=False, keep_links=False)
        try:
            sheets = [sheet_preview(wb[name]) for name in wb.sheetnames]
        finally:
            wb.close()
        after = sha256(path)
        if before != after:
            raise RuntimeError(f"Source changed during preview: {path.name}")
        workbooks.append({"file": path.name, "sha256": before, "sheets": sheets})
        total["workbooks"] += 1
        total["sheets"] += len(sheets)
        for sheet in sheets:
            for key in ("rows", "blank", "duplicate", "review"):
                total[key] += sheet[key]
    return {
        "mode": "PREVIEW_ONLY",
        "publish_enabled": False,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source": str(source),
        "summary": dict(total),
        "workbooks": workbooks,
        "notice": "No source file changed. No row published to dashboard.",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Read-only Excel import preview; never publishes data")
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    report = preview(args.source)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(report["summary"]))


if __name__ == "__main__":
    main()
